"""
Build an Excel workbook of one guide's guidelines in the FSKG statements layout.

The layout mirrors the hand-extracted FSKG statements workbook (README,
Statements, Sample Relations) so the platform's rules can be compared column for
column with an extraction made outside it, and adds two sheets that carry what
the platform actually holds: every field of every rule, and the guide record.

Pure: a guide payload and its rule payloads in, ``.xlsx`` bytes out. Nothing
here talks to the catalog, which is what lets the export job and its tests
share it.
"""

from __future__ import annotations

import io
import json
import math
import re
from collections import Counter
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

STATEMENT_COLUMNS = [
    "Statement ID", "Page (printed)", "Page (PDF)", "Section", "Statement text",
    "Statement type", "Dimension", "Action", "Quantity", "Population", "Food entities",
    "Source", "Source type", "Country", "Publication date",
]
RELATION_COLUMNS = [
    "Relation", "Statement A", "Statement A text", "Statement B", "Statement B text", "Note",
]

# The platform keys rules by UUID, so statement IDs are generated per export. A
# prefix other than the reference workbook's "C" keeps the two from being
# joined on ID by mistake.
STATEMENT_ID_PREFIX = "G"

# Every platform guideline is a rule; Knowledge and Definition statements are
# not stored, so there is nothing else to call them.
STATEMENT_TYPE = "Recommendation"

# First match wins, on substrings of the rule's topic labels. A rule with topics
# but no match is Nutrition; a rule without topics gets no dimension at all.
DIMENSION_RULES = (
    ("FoodSafety*", ("safety", "hygiene", "contamination", "storage", "acrylamide", "allergen")),
    ("Sustainable Diets", ("sustainab", "environment", "climate")),
    ("Food Waste", ("waste", "leftover")),
)
DEFAULT_DIMENSION = "Nutrition"

# Wording follows the reference workbook where it has an equivalent.
TARGET_POPULATION_LABELS = {
    "general_population": "general population",
    "infants": "infants",
    "under_5_years": "children under 5",
    "ages_5_to_18": "children and teenagers 5-18",
    "adults": "adults",
    "elderly": "older adults",
    "pregnant_people": "pregnant women",
    "lactating_people": "breastfeeding mothers",
    "other": "other",
}
LIFE_STAGE_LABELS = {
    "pregnancy": "pregnant women",
    "lactation": "breastfeeding mothers",
    "infancy": "infants",
    "early_childhood": "young children",
    "school_age": "school-age children",
    "adolescence": "teenagers",
    "adulthood": "adults",
    "older_adulthood": "older adults",
}
FREQUENCY_LABELS = {
    "per_meal": "per meal",
    "daily": "daily",
    "weekly": "weekly",
    "monthly": "monthly",
    "occasional": "occasionally",
}
FREQUENCY_PERIODS = {"per_meal": "meal", "daily": "day", "weekly": "week", "monthly": "month"}
OPERATOR_SYMBOLS = {"lt": "<", "lte": "<=", "eq": "", "gte": ">=", "gt": ">", "approx": "~"}

COUNTRY_NAMES = {
    "AT": "Austria", "BE": "Belgium", "BG": "Bulgaria", "HR": "Croatia", "CY": "Cyprus",
    "CZ": "Czechia", "DK": "Denmark", "EE": "Estonia", "FI": "Finland", "FR": "France",
    "DE": "Germany", "GR": "Greece", "HU": "Hungary", "IE": "Ireland", "IT": "Italy",
    "LV": "Latvia", "LT": "Lithuania", "LU": "Luxembourg", "MT": "Malta",
    "NL": "Netherlands", "PL": "Poland", "PT": "Portugal", "RO": "Romania",
    "SK": "Slovakia", "SI": "Slovenia", "ES": "Spain", "SE": "Sweden",
    "GB": "United Kingdom", "UK": "United Kingdom", "NO": "Norway", "IS": "Iceland",
    "CH": "Switzerland", "US": "United States", "CA": "Canada", "AU": "Australia",
    "NZ": "New Zealand",
}

# Platform field order for the Platform Guidelines sheet. Fields not listed are
# appended after these, so nothing the catalog returns is dropped; `embedding`
# is the one exception, being 384 floats no reader wants.
PLATFORM_COLUMNS = [
    "Statement ID", "guide_urn", "id", "sequence_no", "page_no", "title", "section_label",
    "rule_text", "action_type", "target_populations", "life_stage", "age_min_months",
    "age_max_months", "frequency", "quantity", "food_groups", "nutrients",
    "health_conditions", "topic", "guideline_type", "audience", "setting",
    "applicable_regions", "guide_region", "source_refs", "page_summary", "notes", "status",
    "review_status", "visibility", "verifier_user_id", "applicability_status",
    "applicability_start_date", "applicability_end_date", "extractor_name",
    "extractor_run_id", "extraction_model", "enrichment_version", "enrichment_confidence",
    "ai_generated_fields", "enhancements", "embedded_at", "creator", "created_at",
    "updated_at",
]
PLATFORM_EXCLUDED = {"embedding"}

GUIDE_COLUMNS = [
    "urn", "id", "title", "short_title", "description", "status", "language", "region",
    "url", "license", "topic", "audience", "target_audiences", "issuing_authority",
    "responsible_ministry", "organization_urn", "document_type", "legal_status",
    "graphical_model", "evidence_basis", "publication_date", "publication_year",
    "page_count", "revision", "identifiers", "tags", "review_status", "visibility",
    "applicability_status", "applicability_start_date", "applicability_end_date", "notes",
    "artifacts", "guidelines exported", "creator", "created_at", "updated_at",
]
# `content` is the guide's full text and `guidelines` its rule UUIDs, which the
# Platform Guidelines sheet already is.
GUIDE_EXCLUDED = {"content", "guidelines"}

# --------------------------------------------------------------------------- #
# Styling, copied from the reference workbook
# --------------------------------------------------------------------------- #

EXCEL_CELL_LIMIT = 32_767

HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F3B57")
HEADER_ALIGNMENT = Alignment(vertical="center", wrap_text=True)
BODY_FONT = Font(name="Arial", size=10)
BAND_FILL = PatternFill("solid", fgColor="FFF5F8FA")
_GREY = Side(style="thin", color="FFDDDDDD")
BODY_BORDER = Border(left=_GREY, right=_GREY, top=_GREY, bottom=_GREY)
TOP = Alignment(vertical="top")
TOP_WRAP = Alignment(vertical="top", wrap_text=True)

STATEMENT_WIDTHS = {
    "A": 10, "B": 12, "C": 10, "D": 26, "E": 62, "F": 13, "G": 11, "H": 9, "I": 22,
    "J": 22, "K": 26, "L": 68.3, "M": 16, "N": 10, "O": 14,
}
STATEMENT_WRAP = {"Section", "Statement text", "Quantity", "Population", "Food entities", "Source"}
RELATION_WIDTHS = {"A": 14, "B": 10, "C": 45, "D": 10, "E": 45, "F": 50}
PLATFORM_WIDTHS = {
    "rule_text": 62, "page_summary": 50, "notes": 50, "source_refs": 40, "quantity": 30,
    "section_label": 26, "title": 26, "guide_urn": 30, "id": 38,
}

README_FONTS = {
    "title": Font(name="Arial", size=16, bold=True),
    "small": Font(name="Arial", size=10),
    "heading": Font(name="Arial", size=12, bold=True),
    "body": Font(name="Arial", size=11),
}
README_CHARS_PER_LINE = {"title": 80, "small": 135, "heading": 120, "body": 120}
README_LINE_HEIGHT = {"title": 21, "small": 13, "heading": 15, "body": 14}
# Rows advanced after each block: title, sources, then heading/body pairs four
# rows apart, the rhythm of the reference README.
README_ADVANCE = {"title": 2, "small": 3, "heading": 1, "body": 3}


# --------------------------------------------------------------------------- #
# Column mapping
# --------------------------------------------------------------------------- #


def humanize(value: Any) -> str:
    return str(value).replace("_", " ").strip()


def unique(values: Iterable[Any]) -> list:
    """Order-preserving de-duplication that drops empties."""
    return [value for value in dict.fromkeys(values) if value]


def _number_or_last(value: Any) -> float:
    return value if isinstance(value, int) else math.inf


def rule_order(rule: dict) -> tuple:
    """Document order: the guide's sequence number, then the page."""
    return (_number_or_last(rule.get("sequence_no")), _number_or_last(rule.get("page_no")))


def _first_source_ref(rule: dict) -> dict:
    refs = rule.get("source_refs") or []
    return refs[0] if refs and isinstance(refs[0], dict) else {}


def page_pdf(rule: dict) -> int | None:
    page = rule.get("page_no")
    return page if isinstance(page, int) else _first_source_ref(rule).get("page_start")


def section(rule: dict) -> str:
    if rule.get("section_label"):
        return rule["section_label"]
    for ref in rule.get("source_refs") or []:
        if isinstance(ref, dict) and ref.get("section_label"):
            return ref["section_label"]
    return rule.get("title") or ""


def dimension(rule: dict) -> str:
    topics = [humanize(topic).lower() for topic in rule.get("topic") or []]
    if not topics:
        return ""
    for label, needles in DIMENSION_RULES:
        if any(needle in topic for topic in topics for needle in needles):
            return label
    return DEFAULT_DIMENSION


def _format_number(value: Any) -> Any:
    return int(value) if isinstance(value, float) and value.is_integer() else value


def quantity(rule: dict) -> str:
    amount = rule.get("quantity")
    frequency = rule.get("frequency")
    if not isinstance(amount, dict):
        return FREQUENCY_LABELS.get(frequency, "")

    period = amount.get("period") or FREQUENCY_PERIODS.get(frequency)
    raw = (amount.get("raw_text") or "").strip()
    if raw:
        # The source wording keeps ranges ("3 to 5 servings") that the
        # structured operator/value pair cannot hold, so it wins when present.
        return f"{raw}/{period}" if period and period.lower() not in raw.lower() else raw

    symbol = OPERATOR_SYMBOLS.get(amount.get("operator"), "")
    text = f"{symbol}{_format_number(amount.get('value'))} {amount.get('unit') or ''}".strip()
    return f"{text}/{period}" if period else text


def age_range(min_months: Any, max_months: Any) -> str:
    """'1-4 years', '65+ years', 'up to 4 years'; months when years would not be exact."""

    def as_years(months: int, inclusive_upper: bool = False) -> int | None:
        if months % 12 == 0:
            return months // 12
        if inclusive_upper and (months + 1) % 12 == 0:
            return (months + 1) // 12 - 1  # 59 months, inclusive, is "up to 4 years"
        return None

    has_min, has_max = isinstance(min_months, int), isinstance(max_months, int)
    if not (has_min or has_max):
        return ""

    low = as_years(min_months) if has_min else None
    high = as_years(max_months, inclusive_upper=True) if has_max else None
    in_years = (
        (not has_min or low is not None)
        and (not has_max or high is not None)
        and not (has_min and has_max and high <= low)
    )
    if in_years:
        low_text, high_text, unit = low, high, "years"
    else:
        low_text, high_text, unit = min_months, max_months, "months"

    if has_min and has_max:
        return f"{low_text}-{high_text} {unit}"
    if has_min:
        return f"{low_text}+ {unit}"
    return f"up to {high_text} {unit}"


def population(rule: dict) -> str:
    labels = unique(
        TARGET_POPULATION_LABELS.get(value, humanize(value))
        for value in rule.get("target_populations") or []
    )
    if not labels:
        labels = unique(
            LIFE_STAGE_LABELS.get(value, humanize(value)) for value in rule.get("life_stage") or []
        )
    ages = age_range(rule.get("age_min_months"), rule.get("age_max_months"))
    text = "; ".join(labels)
    if ages:
        text = f"{text} ({ages})" if text else ages
    return text


def food_entities(rule: dict) -> str:
    groups = [humanize(group) for group in rule.get("food_groups") or [] if group != "none"]
    nutrients = [str(nutrient).strip() for nutrient in rule.get("nutrients") or []]
    return ", ".join(unique(groups + nutrients))


def guide_year(guide: dict) -> int | None:
    if isinstance(guide.get("publication_year"), int):
        return guide["publication_year"]
    date = str(guide.get("publication_date") or "")
    return int(date[:4]) if date[:4].isdigit() else None


def source_label(guide: dict) -> str:
    """'<authority> - <title> (<year>)', or '(<first year>, rev. <year>)' for a revision."""
    title = guide.get("title") or guide.get("short_title") or guide.get("urn") or ""
    authority = guide.get("issuing_authority")
    label = title if not authority or authority.lower() in title.lower() else f"{authority} - {title}"

    year = guide_year(guide)
    revision = guide.get("revision") if isinstance(guide.get("revision"), dict) else {}
    previous = revision.get("previous_publication_year")
    if year and previous and previous != year:
        return f"{label} ({previous}, rev. {year})"
    return f"{label} ({year})" if year else label


def source_type(guide: dict) -> str:
    value = guide.get("document_type") or ""
    if any(ch.isupper() for ch in value) or " " in value:
        return value  # already human-readable
    return humanize(value.replace("-", "_")).title()


def country(guide: dict, rule: dict | None = None) -> str:
    region = guide.get("region") or (rule or {}).get("guide_region") or ""
    return COUNTRY_NAMES.get(region.upper(), region)


def publication_date(guide: dict) -> str:
    if guide.get("publication_date"):
        return str(guide["publication_date"])[:10]
    return str(guide["publication_year"]) if guide.get("publication_year") else ""


def statement_id(number: int) -> str:
    return f"{STATEMENT_ID_PREFIX}{number:04d}"


def statement_row(number: int, guide: dict, rule: dict) -> dict:
    return {
        "Statement ID": statement_id(number),
        "Page (printed)": None,
        "Page (PDF)": page_pdf(rule),
        "Section": section(rule),
        "Statement text": rule.get("rule_text") or "",
        "Statement type": STATEMENT_TYPE,
        "Dimension": dimension(rule),
        "Action": rule.get("action_type") or "n/a",
        "Quantity": quantity(rule),
        "Population": population(rule),
        "Food entities": food_entities(rule),
        "Source": source_label(guide),
        "Source type": source_type(guide),
        "Country": country(guide, rule),
        "Publication date": publication_date(guide),
    }


def export_filename(guide: dict) -> str:
    """``FSKG_Platform_<guide slug>.xlsx``, safe for a Content-Disposition header."""
    slug = str(guide.get("urn") or "guide").rsplit(":", 1)[-1]
    slug = re.sub(r"[^A-Za-z0-9._-]+", "-", slug).strip("-.") or "guide"
    return f"FSKG_Platform_{slug[:120]}.xlsx"


# --------------------------------------------------------------------------- #
# Writing
# --------------------------------------------------------------------------- #


def excel_value(value: Any) -> Any:
    """A value openpyxl can store: numbers as numbers, lists joined, objects as JSON."""
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, list) and all(isinstance(item, str) for item in value):
        text = "; ".join(value)
    elif isinstance(value, (list, dict)):
        text = json.dumps(value, ensure_ascii=False)
    else:
        text = str(value)
    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    if len(text) > EXCEL_CELL_LIMIT:
        text = text[: EXCEL_CELL_LIMIT - 1] + "…"
    return text or None


def _put(ws, row: int, column: int, value: Any):
    cell = ws.cell(row=row, column=column, value=excel_value(value))
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"  # source text that only looks like a formula
    return cell


def _write_table(
    ws,
    columns: list[str],
    rows: list[dict],
    *,
    widths: dict | None = None,
    wrap: Iterable[str] = (),
    header_height: float = 30,
    row_height: float | None = 35.1,
    default_width: float = 16,
) -> None:
    wrap = set(wrap)
    for index, header in enumerate(columns, 1):
        cell = ws.cell(row=1, column=index, value=header)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, HEADER_ALIGNMENT
    ws.row_dimensions[1].height = header_height

    for row_index, row in enumerate(rows, 2):
        banded = row_index % 2 == 0
        for index, header in enumerate(columns, 1):
            cell = _put(ws, row_index, index, row.get(header))
            cell.font, cell.border = BODY_FONT, BODY_BORDER
            cell.alignment = TOP_WRAP if header in wrap else TOP
            if banded:
                cell.fill = BAND_FILL
        if row_height:
            ws.row_dimensions[row_index].height = row_height

    widths = widths or {}
    for index, header in enumerate(columns, 1):
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = widths.get(letter, widths.get(header, default_width))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


def _counts(values: Iterable[Any]) -> str:
    return ", ".join(f"{value} {count}" for value, count in Counter(values).most_common())


def _readme_blocks(guide: dict, statements: list[dict], rules: list[dict], generated_at: str, origin: str):
    """(style, text) blocks of the README sheet."""
    total = len(statements)

    def filled(column: str) -> int:
        return sum(1 for row in statements if row[column] not in (None, ""))

    where = country(guide) or "unknown region"
    pages = f", {guide['page_count']} pages" if guide.get("page_count") else ""
    span = f" ({statement_id(1)}-{statement_id(total)})" if total else ""

    yield "title", (
        "FSKG (FoodScholar Knowledge Graph) - WiseFood platform export: "
        f"{guide.get('short_title') or guide.get('title') or guide.get('urn')}"
    )
    yield "small", f"Source (1 document, {where}): {source_label(guide)}{pages} [{guide.get('urn')}]."
    yield "heading", "What this is"
    yield "body", (
        "An export of the dietary guidelines (rules) stored on the WiseFood platform for the "
        "guide above, laid out like the hand-extracted FSKG statements workbook so the two can "
        "be compared column for column. Each row of the Statements sheet is one platform "
        "guideline, in the guide's sequence order. Nothing was re-extracted or edited: every "
        "value comes from the platform record, and fields the platform does not hold are left "
        "blank (see Column mapping). Draft and unreviewed rules are included.\n\n"
        f"Generated {generated_at} by {origin}."
    )
    yield "heading", "Result"
    yield "body", (
        f"{total} statements{span}.\n\n"
        f"Action: {_counts(row['Action'] for row in statements) or 'none'}. "
        f"Dimension: {_counts(row['Dimension'] or '(blank)' for row in statements) or 'none'}. "
        f"Filled in: Page (PDF) {filled('Page (PDF)')}/{total}, Section {filled('Section')}/{total}, "
        f"Quantity {filled('Quantity')}/{total}, Population {filled('Population')}/{total}, "
        f"Food entities {filled('Food entities')}/{total}.\n\n"
        f"Status: {_counts(r.get('status') or 'unset' for r in rules) or 'none'}. "
        f"Review status: {_counts(r.get('review_status') or 'unset' for r in rules) or 'none'}. "
        f"Visibility: {_counts(r.get('visibility') or 'unset' for r in rules) or 'none'}. "
        f"Extraction model: {_counts(r.get('extraction_model') or 'unrecorded' for r in rules) or 'none'}."
    )
    yield "heading", "Column mapping"
    yield "body", "\n".join([
        f"Statement ID - generated for this export ({statement_id(1)}, ...); the platform's "
        "guideline UUID is in the Platform Guidelines sheet under the same Statement ID.",
        "Page (printed) - blank: the platform stores the page in the PDF only.",
        "Page (PDF) - page_no, falling back to the first source reference's page_start.",
        "Section - section_label, falling back to a source reference's section_label, then title.",
        "Statement text - rule_text, unchanged.",
        f"Statement type - '{STATEMENT_TYPE}' for every row: the platform stores rules only.",
        "Dimension - derived from the topic labels: food-safety terms give FoodSafety*, "
        "sustainability terms Sustainable Diets, food-waste terms Food Waste, any other topic "
        f"{DEFAULT_DIMENSION}; blank when the rule has no topic labels.",
        "Action - action_type as stored. The platform's vocabulary (eat, drink, use, do, avoid, "
        "prepare, limit, choose, increase, reduce) is wider than the reference workbook's.",
        "Quantity - quantity.raw_text, else operator/value/unit/period, else the frequency.",
        "Population - target_populations (else life_stage), with age_min_months/age_max_months "
        "as an age range.",
        "Food entities - food_groups and nutrients.",
        "Source, Source type, Country, Publication date - from the guide: issuing_authority, "
        "title and publication year; document_type; region; publication_date.",
    ])
    yield "heading", "Not available on the platform"
    yield "body", (
        "Printed page numbers; Knowledge and Definition statements (the platform holds rules, "
        "i.e. recommendations, only); and statement-to-statement relations (equivalentTo, "
        "stricterThan, supports, elaborates, contradicts). The Sample Relations sheet keeps its "
        "headers so the layout matches the reference workbook, and is empty."
    )
    yield "heading", "Sheets"
    yield "body", (
        f"Statements - all {total} platform guidelines in the FSKG columns. Sample Relations - "
        "headers only (see above). Platform Guidelines - every platform field of every rule, keyed "
        "by Statement ID: UUID, sequence number, facets, source references, page summary, review "
        "and visibility state, extraction provenance (lists joined with '; ', objects as JSON). "
        "Guide - the guide record."
    )


def _write_readme(ws, blocks) -> None:
    ws.column_dimensions["A"].width = 100
    row = 1
    for style, text in blocks:
        cell = _put(ws, row, 1, text)
        cell.font = README_FONTS[style]
        cell.alignment = TOP_WRAP
        per_line = README_CHARS_PER_LINE[style]
        lines = sum(max(1, math.ceil(len(part) / per_line)) for part in text.split("\n"))
        ws.row_dimensions[row].height = max(15, lines * README_LINE_HEIGHT[style] + 4)
        row += README_ADVANCE[style]


def _guide_record(guide: dict, exported: int) -> dict:
    record = {key: value for key, value in guide.items() if key not in GUIDE_EXCLUDED}
    record["guidelines exported"] = exported
    record["artifacts"] = [
        {key: artifact.get(key) for key in ("id", "title", "file_type", "language")}
        for artifact in guide.get("artifacts") or []
        if isinstance(artifact, dict)
    ]
    return record


def _columns(preferred: list[str], rows: list[dict], excluded: set[str]) -> list[str]:
    columns = list(preferred)
    for row in rows:
        columns.extend(key for key in row if key not in columns and key not in excluded)
    return columns


def build_guide_workbook(
    guide: dict,
    rules: list[dict],
    *,
    generated_at: str,
    origin: str = "the WiseFood console guide export",
) -> bytes:
    """Return the ``.xlsx`` bytes for one guide and its rules."""
    ordered = sorted((rule for rule in rules if isinstance(rule, dict)), key=rule_order)
    statements = [statement_row(number, guide, rule) for number, rule in enumerate(ordered, 1)]
    platform_rows = [
        {"Statement ID": statement_id(number), **rule} for number, rule in enumerate(ordered, 1)
    ]

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    _write_readme(readme, _readme_blocks(guide, statements, ordered, generated_at, origin))

    _write_table(wb.create_sheet("Statements"), STATEMENT_COLUMNS, statements,
                 widths=STATEMENT_WIDTHS, wrap=STATEMENT_WRAP)
    _write_table(wb.create_sheet("Sample Relations"), RELATION_COLUMNS, [],
                 widths=RELATION_WIDTHS, wrap=RELATION_COLUMNS, header_height=19.5)
    _write_table(wb.create_sheet("Platform Guidelines"),
                 _columns(PLATFORM_COLUMNS, platform_rows, PLATFORM_EXCLUDED), platform_rows,
                 widths=PLATFORM_WIDTHS, wrap={"rule_text", "page_summary", "notes"})
    record = _guide_record(guide, len(ordered))
    _write_table(wb.create_sheet("Guide"), _columns(GUIDE_COLUMNS, [record], GUIDE_EXCLUDED),
                 [record], widths={"title": 40, "urn": 40}, wrap={"title", "description"})

    wb.properties.title = f"FSKG platform export ({guide.get('urn')})"
    wb.properties.creator = "WiseFood FoodScholar guide export"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
