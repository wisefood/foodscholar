"""The console's guide Excel export: the workbook, the job lifecycle and the routes."""

import io
import unittest

from openpyxl import load_workbook

from services.guide_export_jobs import (
    GuideExportError,
    GuideExportJobService,
    GuideExportNotFound,
    GuideExportNotReady,
    normalize_guide_urn,
)
from services.guide_export_workbook import (
    RELATION_COLUMNS,
    STATEMENT_COLUMNS,
    age_range,
    build_guide_workbook,
    export_filename,
    population,
    quantity,
    source_label,
)
from workers.guide_export_worker import BackgroundGuideExportWorker

GUIDE_URN = "urn:guide:ie-healthy-eating"


class _FakeRedisClient:
    def __init__(self):
        self.values = {}
        self.queues = {}
        self.ttl = {}

    def ping(self):
        return True

    def get(self, key):
        return self.values.get(key)

    def set(self, key, value, nx=False, ex=None):
        if nx and key in self.values:
            return False
        self.values[key] = value
        self.ttl[key] = ex
        return True

    def delete(self, key):
        return 1 if self.values.pop(key, None) is not None else 0

    def exists(self, key):
        return 1 if key in self.values else 0

    def rpush(self, key, value):
        self.queues.setdefault(key, []).append(value)
        return len(self.queues[key])

    def blpop(self, key, timeout=0):
        queue = self.queues.get(key, [])
        return (key, queue.pop(0)) if queue else None

    def llen(self, key):
        return len(self.queues.get(key, []))


class _FakeRedis:
    def __init__(self):
        self.client = _FakeRedisClient()


class _FakeGuidelines:
    def __init__(self, rules, error=None):
        self._rules = rules
        self._error = error

    def fetch_all(self):
        if self._error:
            raise self._error
        return list(self._rules)


class _FakeGuide:
    def __init__(self, data, rules, error=None):
        self.data = data
        self.guidelines = _FakeGuidelines(rules, error)


class _FakeClient:
    def __init__(self, guide):
        self.guides = self
        self._guide = guide
        self.requested = []

    def get(self, urn):
        self.requested.append(urn)
        return self._guide


class _FakePool:
    def __init__(self, guide):
        self.client = _FakeClient(guide)
        self.returned = 0

    def get_client(self):
        return self.client

    def return_client(self, client):
        self.returned += 1


GUIDE = {
    "urn": GUIDE_URN,
    "title": "Healthy Eating, Food Safety and Food Legislation",
    "issuing_authority": "FSAI",
    "region": "IE",
    "document_type": "national_dietary_guide",
    "publication_year": 2021,
    "publication_date": "2021-03-01T00:00:00Z",
    "revision": {"previous_publication_year": 2019},
    "page_count": 90,
    "content": "the whole guide text",
    "guidelines": ["a", "b"],
}


def _rule(sequence_no, text, **fields):
    return {
        "id": f"id-{sequence_no}",
        "guide_urn": GUIDE_URN,
        "sequence_no": sequence_no,
        "rule_text": text,
        "action_type": "eat",
        **fields,
    }


RULES = [
    _rule(2, "Eat oily fish once a week.", page_no=4, section_label="At a glance",
          quantity={"operator": "eq", "value": 1.0, "unit": "serving", "period": "week",
                    "raw_text": None},
          target_populations=["general_population"], food_groups=["protein_foods"],
          nutrients=["vitamin D"], topic=["healthy_eating"], embedding=[0.1] * 384),
    _rule(1, "=SUM(A1) is text, not a formula.\x07", page_no=None,
          source_refs=[{"page_start": 3, "section_label": "Introduction"}],
          topic=["food_safety"], status="draft", review_status="unreviewed",
          enhancements=[{"blob": "y" * 40_000}], unexpected_new_field={"nested": 1}),
    _rule(None, "Children under five should not be given whole nuts.", page_no=12,
          title="Choking risks", life_stage=["early_childhood"], age_min_months=12,
          age_max_months=59, frequency="daily"),
]


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        content = build_guide_workbook(GUIDE, RULES, generated_at="2026-10-08T12:00:00Z")
        self.wb = load_workbook(io.BytesIO(content))

    def rows(self, sheet):
        ws = self.wb[sheet]
        header = [cell.value for cell in ws[1]]
        return header, [dict(zip(header, (c.value for c in row))) for row in ws.iter_rows(min_row=2)]

    def test_the_reference_layout_comes_first(self):
        self.assertEqual(
            self.wb.sheetnames, ["README", "Statements", "Sample Relations", "Platform Guidelines", "Guide"]
        )
        self.assertEqual(self.rows("Statements")[0], STATEMENT_COLUMNS)
        self.assertEqual(self.rows("Sample Relations"), (RELATION_COLUMNS, []))

    def test_rules_follow_the_guide_sequence_with_unnumbered_rules_last(self):
        _, statements = self.rows("Statements")
        self.assertEqual([row["Statement ID"] for row in statements], ["G0001", "G0002", "G0003"])
        self.assertTrue(statements[0]["Statement text"].startswith("=SUM(A1) is text"))
        self.assertEqual(statements[2]["Statement text"], RULES[2]["rule_text"])

    def test_text_that_looks_like_a_formula_stays_text(self):
        cell = self.wb["Statements"]["E2"]
        self.assertEqual(cell.data_type, "s")
        self.assertNotIn("\x07", cell.value)

    def test_columns_are_mapped_from_the_platform_fields(self):
        _, statements = self.rows("Statements")
        first, second, third = statements
        self.assertEqual((first["Page (PDF)"], first["Section"], first["Dimension"]),
                         (3, "Introduction", "FoodSafety*"))
        self.assertEqual(second["Quantity"], "1 serving/week")
        self.assertEqual(second["Population"], "general population")
        self.assertEqual(second["Food entities"], "protein foods, vitamin D")
        self.assertEqual(second["Dimension"], "Nutrition")
        self.assertEqual((third["Section"], third["Dimension"], third["Quantity"]),
                         ("Choking risks", None, "daily"))
        self.assertEqual(third["Population"], "young children (1-4 years)")
        self.assertIsNone(third["Page (printed)"])
        self.assertEqual(first["Statement type"], "Recommendation")
        self.assertEqual(first["Source"], "FSAI - Healthy Eating, Food Safety and Food Legislation (2019, rev. 2021)")
        self.assertEqual((first["Source type"], first["Country"], first["Publication date"]),
                         ("National Dietary Guide", "Ireland", "2021-03-01"))

    def test_every_platform_field_is_kept_except_the_embedding(self):
        header, rows = self.rows("Platform Guidelines")
        self.assertIn("unexpected_new_field", header)
        self.assertNotIn("embedding", header)
        self.assertEqual(len(rows[0]["enhancements"]), 32_767)
        self.assertEqual(rows[0]["status"], "draft")

    def test_the_guide_sheet_leaves_out_the_full_text(self):
        header, rows = self.rows("Guide")
        self.assertNotIn("content", header)
        self.assertNotIn("guidelines", header)
        self.assertEqual(rows[0]["guidelines exported"], 3)

    def test_an_empty_guide_still_builds(self):
        wb = load_workbook(io.BytesIO(build_guide_workbook(GUIDE, [], generated_at="now")))
        self.assertEqual(wb["Statements"].max_row, 1)


class MappingTests(unittest.TestCase):
    def test_age_ranges(self):
        self.assertEqual(age_range(12, 59), "1-4 years")
        self.assertEqual(age_range(780, None), "65+ years")
        self.assertEqual(age_range(None, 59), "up to 4 years")
        self.assertEqual(age_range(0, 11), "0-11 months")
        self.assertEqual(age_range(6, 12), "6-12 months")
        self.assertEqual(age_range(None, None), "")

    def test_quantity_prefers_the_source_wording(self):
        rule = {"quantity": {"operator": "gte", "value": 3, "unit": "servings",
                             "period": "day", "raw_text": "3 to 5 servings"}}
        self.assertEqual(quantity(rule), "3 to 5 servings/day")
        rule["quantity"]["raw_text"] = None
        self.assertEqual(quantity(rule), ">=3 servings/day")
        self.assertEqual(quantity({"frequency": "weekly"}), "weekly")
        self.assertEqual(quantity({}), "")

    def test_population_falls_back_to_life_stage(self):
        self.assertEqual(population({"life_stage": ["pregnancy"]}), "pregnant women")
        self.assertEqual(population({}), "")

    def test_source_label_without_revision_or_authority(self):
        self.assertEqual(source_label({"title": "Guide", "publication_year": 2020}), "Guide (2020)")

    def test_filename_is_header_safe(self):
        self.assertEqual(export_filename({"urn": GUIDE_URN}), "FSKG_Platform_ie-healthy-eating.xlsx")
        self.assertEqual(export_filename({"urn": 'urn:guide:a"b\r\nc'}), "FSKG_Platform_a-b-c.xlsx")


class JobTests(unittest.TestCase):
    def setUp(self):
        self.redis = _FakeRedis()
        self.pool = _FakePool(_FakeGuide(dict(GUIDE), RULES))
        self.service = GuideExportJobService(redis_client=self.redis, platform_pool=self.pool)
        self.worker = BackgroundGuideExportWorker(job_service=self.service)

    def drain(self):
        while (job := self.service.pop_next_job(timeout=0)) is not None:
            self.worker._process_job(job)

    def test_a_slug_and_a_urn_name_the_same_guide(self):
        self.assertEqual(normalize_guide_urn("ie-healthy-eating"), GUIDE_URN)
        self.assertEqual(normalize_guide_urn(GUIDE_URN), GUIDE_URN)
        for bad in ("", "  ", "a b", "a/b"):
            with self.assertRaises(GuideExportError):
                normalize_guide_urn(bad)

    def test_nothing_exported_yet(self):
        self.assertEqual(self.service.get_job_response(GUIDE_URN).status, "not_found")
        with self.assertRaises(GuideExportNotFound):
            self.service.get_file(GUIDE_URN)

    def test_repeated_requests_while_queued_share_one_job(self):
        first = self.service.enqueue_job("ie-healthy-eating")
        second = self.service.enqueue_job(GUIDE_URN)
        self.assertEqual(first["job_id"], second["job_id"])
        self.assertEqual(self.redis.client.llen(self.service.queue_key), 1)
        with self.assertRaises(GuideExportNotReady):
            self.service.get_file(GUIDE_URN)

        forced = self.service.enqueue_job(GUIDE_URN, force=True)
        self.assertNotEqual(forced["job_id"], first["job_id"])

    def test_an_export_runs_to_a_downloadable_workbook(self):
        self.service.enqueue_job(GUIDE_URN)
        self.drain()

        response = self.service.get_job_response(GUIDE_URN)
        self.assertEqual((response.status, response.stage), ("succeeded", "done"))
        self.assertTrue(response.download_ready)
        self.assertEqual(response.total_guidelines, 3)
        self.assertEqual(response.filename, "FSKG_Platform_ie-healthy-eating.xlsx")

        filename, content = self.service.get_file(GUIDE_URN)
        self.assertEqual(len(content), response.size_bytes)
        self.assertEqual(load_workbook(io.BytesIO(content))["Statements"].max_row, 4)
        self.assertEqual(self.pool.client.requested, [GUIDE_URN])
        self.assertEqual(self.pool.returned, 1)
        # The lock is released and both keys expire.
        self.assertFalse(self.service.has_lock(GUIDE_URN))
        file_key = self.service._file_key(response.job_id)
        self.assertEqual(self.redis.client.ttl[file_key], self.service.ttl_seconds)
        self.assertEqual(self.redis.client.ttl[self.service._job_key(GUIDE_URN)], self.service.ttl_seconds)

    def test_a_finished_export_is_redone_when_asked_again(self):
        self.service.enqueue_job(GUIDE_URN)
        self.drain()
        done = self.service.get_job_response(GUIDE_URN).job_id
        again = self.service.enqueue_job(GUIDE_URN)
        self.assertNotEqual(again["job_id"], done)
        self.assertEqual(again["status"], "queued")

    def test_an_expired_file_is_reported_as_gone(self):
        self.service.enqueue_job(GUIDE_URN)
        self.drain()
        job_id = self.service.get_job_response(GUIDE_URN).job_id
        self.redis.client.delete(self.service._file_key(job_id))
        self.assertFalse(self.service.get_job_response(GUIDE_URN).download_ready)
        with self.assertRaises(GuideExportNotFound):
            self.service.get_file(GUIDE_URN)

    def test_a_failure_is_recorded_and_the_lock_released(self):
        self.pool = _FakePool(_FakeGuide(dict(GUIDE), [], error=RuntimeError("catalog said 502")))
        self.service._platform_pool = self.pool
        self.service.enqueue_job(GUIDE_URN)
        self.drain()

        response = self.service.get_job_response(GUIDE_URN)
        self.assertEqual(response.status, "failed")
        self.assertIn("502", response.error)
        self.assertFalse(response.download_ready)
        self.assertFalse(self.service.has_lock(GUIDE_URN))
        self.assertEqual(self.pool.returned, 1)

    def test_a_superseded_job_is_skipped(self):
        stale = self.service.enqueue_job(GUIDE_URN)
        self.service.enqueue_job(GUIDE_URN, force=True)
        self.service.pop_next_job(timeout=0)  # the stale queue entry
        self.worker._process_job(stale)
        self.assertEqual(self.service.get_job_response(GUIDE_URN).status, "queued")

    def test_a_job_whose_worker_died_is_stalled_and_can_be_requeued(self):
        job = self.service.enqueue_job(GUIDE_URN)
        self.service._update(GUIDE_URN, status="running")  # and no lock is held
        self.assertEqual(self.service.get_job_response(GUIDE_URN).status, "stalled")
        self.assertNotEqual(self.service.enqueue_job(GUIDE_URN)["job_id"], job["job_id"])


class RouteTests(unittest.TestCase):
    def setUp(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient

        from api.v1 import guidelines

        self.service = GuideExportJobService(
            redis_client=_FakeRedis(), platform_pool=_FakePool(_FakeGuide(dict(GUIDE), RULES))
        )
        self._original = guidelines.export_service
        guidelines.export_service = self.service
        app = FastAPI()
        app.include_router(guidelines.router, prefix="/api/v1")
        self.client = TestClient(app)

    def tearDown(self):
        from api.v1 import guidelines

        guidelines.export_service = self._original

    def test_queue_poll_and_download(self):
        queued = self.client.post(f"/api/v1/guidelines/export/{GUIDE_URN}")
        self.assertEqual(queued.status_code, 202)
        self.assertEqual(queued.json()["status"], "queued")

        self.assertEqual(self.client.get(f"/api/v1/guidelines/export/{GUIDE_URN}/download").status_code, 409)

        worker = BackgroundGuideExportWorker(job_service=self.service)
        worker._process_job(self.service.pop_next_job(timeout=0))

        status = self.client.get(f"/api/v1/guidelines/export/{GUIDE_URN}").json()
        self.assertTrue(status["download_ready"])

        download = self.client.get(f"/api/v1/guidelines/export/{GUIDE_URN}/download")
        self.assertEqual(download.status_code, 200)
        self.assertTrue(download.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
        self.assertIn('filename="FSKG_Platform_ie-healthy-eating.xlsx"',
                      download.headers["content-disposition"])
        self.assertEqual(load_workbook(io.BytesIO(download.content)).sheetnames[1], "Statements")

    def test_unknown_and_invalid_guides(self):
        self.assertEqual(self.client.get("/api/v1/guidelines/export/urn:guide:nothing").json()["status"],
                         "not_found")
        self.assertEqual(self.client.get("/api/v1/guidelines/export/urn:guide:nothing/download").status_code,
                         404)
        self.assertEqual(self.client.post("/api/v1/guidelines/export/a%20b").status_code, 400)


if __name__ == "__main__":
    unittest.main()
